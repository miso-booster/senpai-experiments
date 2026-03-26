from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


def _safe_frame(df: pd.DataFrame | None) -> pd.DataFrame:
    return df if df is not None else pd.DataFrame()


def _filter_arch(df: pd.DataFrame, arch_id: str) -> pd.DataFrame:
    if df.empty or "arch_id" not in df.columns:
        return pd.DataFrame()
    return df[df["arch_id"] == arch_id].copy()


def _review_instructions() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "step": [1, 2, 3, 4],
            "instruction": [
                "Review the 'review_assignments' sheet only.",
                "If the proposed cluster is correct, leave final_cluster_id unchanged and set reviewed=TRUE.",
                "If it is wrong, change final_cluster_id to another cluster id such as C2, or create a new one such as C4, then set reviewed=TRUE.",
                "After review, run: python -m app.rq3_clustering.review_metrics",
            ],
        }
    )


def _apply_review_validation(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    if "review_assignments" not in workbook.sheetnames:
        workbook.save(workbook_path)
        return

    sheet = workbook["review_assignments"]
    header_map = {cell.value: cell.column for cell in sheet[1] if cell.value}
    reviewed_column = header_map.get("reviewed")
    if reviewed_column is not None and sheet.max_row >= 2:
        column_letter = sheet.cell(row=1, column=reviewed_column).column_letter
        validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        validation.prompt = "Set TRUE when this row has been reviewed."
        validation.error = "Use TRUE or FALSE."
        sheet.add_data_validation(validation)
        validation.add(f"{column_letter}2:{column_letter}{sheet.max_row}")

    sheet.freeze_panes = "A2"
    workbook.save(workbook_path)


def save_outputs(
    out_dir: Path,
    architecture_tables: dict[str, pd.DataFrame],
    inventory: pd.DataFrame,
    duplicates: pd.DataFrame,
    *,
    overall_summary: pd.DataFrame | None,
    prompt_summary: pd.DataFrame | None,
    cluster_summary: pd.DataFrame | None,
    cluster_assignments: pd.DataFrame | None,
    review_assignments: pd.DataFrame | None,
    review_cluster_reference: pd.DataFrame | None,
    export_csvs: bool,
) -> list[Path]:
    """Write the phase-1 LLM clustering outputs plus the human review workbook."""
    out_dir.mkdir(parents=True, exist_ok=True)

    overall_summary = _safe_frame(overall_summary)
    prompt_summary = _safe_frame(prompt_summary)
    cluster_summary = _safe_frame(cluster_summary)
    cluster_assignments = _safe_frame(cluster_assignments)
    review_assignments = _safe_frame(review_assignments)
    review_cluster_reference = _safe_frame(review_cluster_reference)

    written_paths: list[Path] = []
    for arch_id, table in architecture_tables.items():
        workbook_path = out_dir / f"{arch_id}_answer_clustering.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            table.to_excel(writer, sheet_name="answers_matrix")
            arch_inventory = inventory[inventory["arch_id"] == arch_id].copy()
            arch_inventory.to_excel(writer, index=False, sheet_name="inventory")

            arch_prompt_summary = _filter_arch(prompt_summary, arch_id)
            arch_prompt_summary.to_excel(writer, index=False, sheet_name="prompt_summary")

            arch_cluster_summary = _filter_arch(cluster_summary, arch_id)
            arch_cluster_summary.to_excel(writer, index=False, sheet_name="cluster_summary")

            arch_assignments = _filter_arch(cluster_assignments, arch_id)
            arch_assignments.to_excel(writer, index=False, sheet_name="cluster_assignments")

            arch_review_assignments = _filter_arch(review_assignments, arch_id)
            arch_review_assignments.to_excel(writer, index=False, sheet_name="review_assignments")

            arch_review_reference = _filter_arch(review_cluster_reference, arch_id)
            arch_review_reference.to_excel(writer, index=False, sheet_name="review_cluster_reference")

            arch_overall_summary = _filter_arch(overall_summary, arch_id)
            arch_overall_summary.to_excel(writer, index=False, sheet_name="overall_summary")

            if not duplicates.empty:
                arch_duplicates = duplicates[duplicates["arch_id"] == arch_id].copy()
                if not arch_duplicates.empty:
                    arch_duplicates.to_excel(writer, index=False, sheet_name="duplicates")
        written_paths.append(workbook_path)
        _apply_review_validation(workbook_path)

        if export_csvs:
            table.to_csv(out_dir / f"{arch_id}_answer_clustering.csv", encoding="utf-8")
            arch_inventory.to_csv(out_dir / f"{arch_id}_inventory.csv", index=False, encoding="utf-8")
            arch_prompt_summary.to_csv(out_dir / f"{arch_id}_prompt_summary.csv", index=False, encoding="utf-8")
            arch_cluster_summary.to_csv(out_dir / f"{arch_id}_cluster_summary.csv", index=False, encoding="utf-8")
            arch_assignments.to_csv(out_dir / f"{arch_id}_cluster_assignments.csv", index=False, encoding="utf-8")
            arch_review_assignments.to_csv(
                out_dir / f"{arch_id}_review_assignments.csv",
                index=False,
                encoding="utf-8",
            )
            arch_review_reference.to_csv(
                out_dir / f"{arch_id}_review_cluster_reference.csv",
                index=False,
                encoding="utf-8",
            )
            arch_overall_summary.to_csv(out_dir / f"{arch_id}_overall_summary.csv", index=False, encoding="utf-8")
            if not duplicates.empty:
                arch_duplicates = duplicates[duplicates["arch_id"] == arch_id]
                if not arch_duplicates.empty:
                    arch_duplicates.to_csv(
                        out_dir / f"{arch_id}_duplicates.csv",
                        index=False,
                        encoding="utf-8",
                    )

    manifest_path = out_dir / "RQ3_clustering_manifest.xlsx"
    with pd.ExcelWriter(manifest_path, engine="openpyxl") as writer:
        overall_summary.to_excel(writer, index=False, sheet_name="overall_summary")
        prompt_summary.to_excel(writer, index=False, sheet_name="prompt_summary")
        cluster_summary.to_excel(writer, index=False, sheet_name="cluster_summary")
        cluster_assignments.to_excel(writer, index=False, sheet_name="cluster_assignments")
        review_assignments.to_excel(writer, index=False, sheet_name="review_assignments")
        review_cluster_reference.to_excel(writer, index=False, sheet_name="review_cluster_reference")
        inventory.to_excel(writer, index=False, sheet_name="inventory")
        if not duplicates.empty:
            duplicates.to_excel(writer, index=False, sheet_name="duplicates")
        pd.DataFrame(
            [
                {
                    "arch_id": arch_id,
                    "matrix_rows": len(table),
                    "matrix_columns": len(table.columns),
                    "prompt_count": len(_filter_arch(prompt_summary, arch_id)),
                    "cluster_count": len(_filter_arch(cluster_summary, arch_id)),
                    "assignment_count": len(_filter_arch(cluster_assignments, arch_id)),
                }
                for arch_id, table in architecture_tables.items()
            ]
        ).to_excel(writer, index=False, sheet_name="architectures")
    written_paths.append(manifest_path)
    _apply_review_validation(manifest_path)

    review_workbook_path = out_dir / "RQ3_clustering_review_template.xlsx"
    with pd.ExcelWriter(review_workbook_path, engine="openpyxl") as writer:
        _review_instructions().to_excel(writer, index=False, sheet_name="instructions")
        review_assignments.to_excel(writer, index=False, sheet_name="review_assignments")
        review_cluster_reference.to_excel(writer, index=False, sheet_name="review_cluster_reference")
    written_paths.append(review_workbook_path)
    _apply_review_validation(review_workbook_path)

    if export_csvs:
        overall_summary.to_csv(out_dir / "rq3_clustering_overall_summary.csv", index=False, encoding="utf-8")
        prompt_summary.to_csv(out_dir / "rq3_clustering_prompt_summary.csv", index=False, encoding="utf-8")
        cluster_summary.to_csv(out_dir / "rq3_clustering_cluster_summary.csv", index=False, encoding="utf-8")
        cluster_assignments.to_csv(
            out_dir / "rq3_clustering_cluster_assignments.csv",
            index=False,
            encoding="utf-8",
        )
        review_assignments.to_csv(
            out_dir / "rq3_clustering_review_assignments.csv",
            index=False,
            encoding="utf-8",
        )
        review_cluster_reference.to_csv(
            out_dir / "rq3_clustering_review_cluster_reference.csv",
            index=False,
            encoding="utf-8",
        )

    return written_paths


def save_reviewed_outputs(
    out_dir: Path,
    reviewed_assignments: pd.DataFrame,
    reviewed_overall_summary: pd.DataFrame,
    reviewed_prompt_summary: pd.DataFrame,
    reviewed_cluster_summary: pd.DataFrame,
    reviewed_pairwise_wilcoxon: pd.DataFrame,
    *,
    export_csvs: bool,
) -> list[Path]:
    """Write the phase-2 reviewed clustering metrics."""
    out_dir.mkdir(parents=True, exist_ok=True)

    written_paths: list[Path] = []
    for arch_id in sorted(reviewed_assignments["arch_id"].astype(str).unique().tolist()):
        workbook_path = out_dir / f"{arch_id}_answer_clustering_reviewed.xlsx"
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            _filter_arch(reviewed_overall_summary, arch_id).to_excel(
                writer, index=False, sheet_name="overall_summary"
            )
            _filter_arch(reviewed_prompt_summary, arch_id).to_excel(
                writer, index=False, sheet_name="prompt_summary"
            )
            _filter_arch(reviewed_cluster_summary, arch_id).to_excel(
                writer, index=False, sheet_name="cluster_summary"
            )
            _filter_arch(reviewed_assignments, arch_id).to_excel(
                writer, index=False, sheet_name="reviewed_assignments"
            )
        written_paths.append(workbook_path)

    manifest_path = out_dir / "RQ3_clustering_reviewed_metrics.xlsx"
    with pd.ExcelWriter(manifest_path, engine="openpyxl") as writer:
        reviewed_overall_summary.to_excel(writer, index=False, sheet_name="overall_summary")
        reviewed_prompt_summary.to_excel(writer, index=False, sheet_name="prompt_summary")
        reviewed_cluster_summary.to_excel(writer, index=False, sheet_name="cluster_summary")
        reviewed_assignments.to_excel(writer, index=False, sheet_name="reviewed_assignments")
        if not reviewed_pairwise_wilcoxon.empty:
            reviewed_pairwise_wilcoxon.to_excel(writer, index=False, sheet_name="pairwise_wilcoxon")
    written_paths.append(manifest_path)

    if export_csvs:
        reviewed_overall_summary.to_csv(
            out_dir / "rq3_clustering_reviewed_overall_summary.csv",
            index=False,
            encoding="utf-8",
        )
        reviewed_prompt_summary.to_csv(
            out_dir / "rq3_clustering_reviewed_prompt_summary.csv",
            index=False,
            encoding="utf-8",
        )
        reviewed_cluster_summary.to_csv(
            out_dir / "rq3_clustering_reviewed_cluster_summary.csv",
            index=False,
            encoding="utf-8",
        )
        reviewed_assignments.to_csv(
            out_dir / "rq3_clustering_reviewed_assignments.csv",
            index=False,
            encoding="utf-8",
        )
        if not reviewed_pairwise_wilcoxon.empty:
            reviewed_pairwise_wilcoxon.to_csv(
                out_dir / "rq3_clustering_reviewed_pairwise_wilcoxon.csv",
                index=False,
                encoding="utf-8",
            )

    return written_paths
